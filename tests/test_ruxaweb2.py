#!/usr/bin/env python3
# SPDX-FileCopyrightText: 2026 Bertrand Kempf
# SPDX-License-Identifier: EUPL-1.2
"""Tests de bout en bout de src/ruxaweb2.html (RUXAWEB 2, audit par observations ; Playwright, Chromium).

Installation (une fois) :
  python3 -m pip install playwright openpyxl
  python3 -m playwright install chromium
Lancement :
  python3 tests/test_ruxaweb2.py
Code retour 0 si tout passe, 1 sinon.
"""
import csv
import io
import json
import pathlib
import sys
import tempfile

import openpyxl
from playwright.sync_api import sync_playwright

ROOT = pathlib.Path(__file__).resolve().parent.parent
URL = (ROOT / 'src' / 'ruxaweb2.html').as_uri()
resultats = []


def verifier(nom, condition, detail=''):
    resultats.append(bool(condition))
    print(('OK    ' if condition else 'ÉCHEC ') + nom + ('' if condition else ' , ' + str(detail)))


def statut(page, ecran, crit):
    """Statut d'un critère sur un écran (par index), lu dans la grille."""
    page.click('[data-action="vue"][data-vue="grille"]')
    ligne = page.locator('tr', has=page.locator('[data-action="critere"][data-id="%s"]' % crit))
    return ligne.locator('td.statut').nth(ecran).inner_text().split('\n')[0].strip()


def observer(page, texte, point=None, element=None, tag='nc', partage=None, crit=None):
    page.click('[data-action="vue"][data-vue="audit"]')
    page.fill('#obs-texte', texte)
    if point:
        page.select_option('#obs-point', point)
    if element is not None:
        page.select_option('#obs-element', str(element))
    if crit:
        page.select_option('#obs-ajout', crit)
    page.click('[data-action="tag"][data-tag="%s"]' % tag)
    if partage:
        nom, quels = partage
        if nom:
            page.fill('#obs-partage-nom', nom)
        cases = page.locator('[data-ecran-partage]').all()
        for i, cb in enumerate(cases):
            if quels is None or i in quels:
                cb.check()
    page.click('#form-obs button[type=submit]')


def ecran(page, i):
    page.click('[data-action="vue"][data-vue="audit"]')
    page.locator('.ecrans [data-action="ecran"]').nth(i).click()


def telecharger(page, action):
    with page.expect_download() as d:
        page.click('#btn-menu')
        page.click('[data-action="%s"]' % action)
    chemin = pathlib.Path(tempfile.mkdtemp()) / d.value.suggested_filename
    d.value.save_as(chemin)
    return chemin


def main():
    with sync_playwright() as p:
        nav = p.chromium.launch()
        page = nav.new_page(viewport={'width': 1360, 'height': 950}, accept_downloads=True)
        erreurs = []
        page.on('pageerror', lambda e: erreurs.append(str(e)))
        page.goto(URL)
        verifier('Ouverture : tableau de bord', page.locator('h1').inner_text() == 'Tableau de bord')
        verifier('Aide présente avec le parcours', page.locator('#aide .parcours li').count() >= 9)
        txt = page.content().lower()
        verifier('Aucune mention de Claude', 'claude' not in txt and 'anthropic' not in txt)

        page.fill('#audit-nom', 'Boutique test')
        page.locator('#audit-nom').blur()
        page.click('[data-action="nouvel-onglet"]')
        verifier('Nouvel onglet : 3 écrans par défaut', page.locator('[data-ecran-nom]').count() == 3)
        # Écran 1 : formulaire présent ; les autres sans formulaire.
        page.click('[data-action="contenu"][data-d="formulaire"] >> nth=0')
        sw = page.locator('[data-action="contenu"][data-d="formulaire"]').first
        verifier('Contenu des écrans : interrupteur activé', sw.get_attribute('role') == 'switch' and sw.get_attribute('aria-checked') == 'true')
        page.click('[data-action="vue"][data-vue="audit"] >> nth=-1')
        verifier('Audit : formulaire d\'observation', page.locator('#form-obs').count() == 1)
        verifier('Aide : étape Tour à vide en cours', 'Tour à vide' in page.locator('#aide .ici').inner_text())

        verifier('Configurateur : 11.1 NA sans formulaire', statut(page, 1, 'raweb-11-1') == 'NA')
        verifier('Configurateur : 11.1 NT avec formulaire', statut(page, 0, 'raweb-11-1') == 'NT')

        # Partout #nc : NC sur les écrans cochés.
        observer(page, 'Menu principal : focus invisible.', point='focus-invisible', partage=('Menu principal', None))
        verifier('Élément répété, #nc : NC sur les 3 écrans', [statut(page, i, 'raweb-10-7') for i in range(3)] == ['NC', 'NC', 'NC'])
        # Partout #c : conforme ici, à confirmer ailleurs.
        observer(page, 'Liens du pied de page explicites.', point='nom-non-pertinent', element=2, tag='c', partage=('Pied de page', None))
        verifier('Élément répété, #c : C ici, à confirmer ailleurs', [statut(page, i, 'raweb-6-1') for i in range(3)] == ['C', 'C ?', 'C ?'])
        # Sans écran coché, l'observation reste propre à l'écran ; sur un écran sans formulaire, le critère reste NA.
        observer(page, 'Étiquette du champ e-mail absente.', point='etiquette')
        verifier('Sans autre écran coché : 11.1 NC sur cet écran seulement', [statut(page, i, 'raweb-11-1') for i in range(3)] == ['NC', 'NA', 'NA'])
        observer(page, 'Erreur non restituée.', point='erreur-saisie', partage=(None, None))
        verifier('Aussi sur, sans nom : 11.10 NA là où le formulaire est absent', [statut(page, i, 'raweb-11-10') for i in range(3)] == ['NC', 'NA', 'NA'])
        observer(page, 'Menu : état ouvert non annoncé.', point='role-etat', partage=('Menu', [0]))
        verifier('Aussi sur un seul écran : 7.1 NC sur les deux, NT sur le troisième', [statut(page, i, 'raweb-7-1') for i in range(3)] == ['NC', 'NC', 'NT'])
        page.click('[data-action="vue"][data-vue="audit"]')
        verifier('Éléments répétés : liste unique à gauche', 'éléments répétés' in page.locator('nav.ecrans').inner_text().lower() and 'Menu principal' in page.locator('nav.ecrans').inner_text())

        # Noms d'élément : casse, accents et espaces ne créent pas un nouvel élément.
        ecran(page, 1)
        observer(page, 'Lien du menu sans nom.', point='nom-absent', element=1, partage=('  menu   PRINCIPAL ', []))
        derniere = page.evaluate('RUXAWEB.state().onglets[0].obs.slice(-1)[0]')
        verifier('Élément : nom ramené à l\'orthographe existante', derniere.get('elementNom') == 'Menu principal', derniere.get('elementNom'))
        page.click('[data-action="vue"][data-vue="audit"]')
        verifier('Éléments répétés : pas de doublon', page.locator('nav.ecrans .elt', has_text='Menu principal').count() == 1)
        ecran(page, 0)

        # Note libre : aucun statut.
        avant = page.evaluate('RUXAWEB.state().onglets[0].obs.length')
        observer(page, 'Page chargée lentement.', tag='c')
        verifier('Note libre enregistrée sans critère', page.evaluate('RUXAWEB.state().onglets[0].obs.slice(-1)[0].criteres.length') == 0 and page.evaluate('RUXAWEB.state().onglets[0].obs.length') == avant + 1)
        page.fill('#obs-texte', '')
        page.click('#form-obs button[type=submit]')
        verifier('Observation vide : message', 'Écris l\'observation' in page.locator('#statut').inner_text())

        # Retest (#nt).
        page.click('[data-action="vue"][data-vue="grille"]')
        page.click('[data-action="critere"][data-id="raweb-10-7"]')
        page.locator('#dlg [data-action="retest"]').first.click()
        page.click('[data-action="fermer-dlg"]')
        verifier('Retest : 10.7 repasse NT sur cet écran', statut(page, 0, 'raweb-10-7') == 'NT')
        verifier('Retest : les autres écrans gardent NC', statut(page, 1, 'raweb-10-7') == 'NC')
        ecran(page, 0)
        verifier('Retest : observation barrée', page.locator('.obs.invalide').count() >= 1)

        # Taux : hors sélection (7.1 NC) ne change que le taux élargi.
        oid = page.evaluate('RUXAWEB.state().onglets[0].id')
        page.click('[data-action="vue"][data-vue="cloture"]')
        verifier('Clôture : un critère à la fois, avec sa consigne', page.locator('#clo-titre').count() == 1 and 'Comment vérifier' in page.locator('main').inner_text())
        premier = page.locator('#clo-titre').inner_text()
        page.locator('[data-action="clo-crit"]').first.click()
        verifier('Clôture : passage à un autre critère', page.locator('#clo-titre').inner_text() != premier)
        tous = page.locator('[data-action="clo-tous"]')
        if tous.count():
            cid = tous.get_attribute('data-crit')
            tous.click()
            verifier('Clôture : conforme sur les écrans restants en un clic', all(statut(page, i, cid) in ('C', 'NA') for i in range(3)))
            page.click('[data-action="vue"][data-vue="cloture"]')
        for _ in range(400):
            b = page.locator('[data-action="clo"][data-tag="c"]')
            if b.count() == 0:
                break
            b.first.click()
        sel = page.evaluate('RUXAWEB.taux("%s")' % oid)
        ela = page.evaluate('RUXAWEB.taux("%s", true)' % oid)
        verifier('Clôture : plus rien à trancher', sel['complet'] and ela['complet'], (sel, ela))
        verifier('Taux sélection SIP : 5 NC (6.2, 10.7, 11.1, 11.2, 11.10), sans 7.1', sel['C'] == sel['total'] - 5, sel)
        verifier('Taux élargi : 7.1 en plus', ela['C'] == ela['total'] - 6, ela)
        page.click('[data-action="valider"]')
        verifier('Validation : résumé', page.locator('h1').inner_text().startswith('Résumé'))
        verifier('Résumé : deux taux', 'Taux sélection SIP' in page.locator('main').inner_text() and 'Taux élargi' in page.locator('main').inner_text())

        # Exports.
        x = telecharger(page, 'export-xlsx')
        ws = openpyxl.load_workbook(x).active
        lignes = list(ws.iter_rows(values_only=True))
        verifier('Excel : une ligne par critère', len(lignes) == 1 + 61, len(lignes))
        verifier('Excel : numéro gardé en texte', any(r[1] == '10.10' for r in lignes))
        verifier('Excel : 10.7 NC', any(r[1] == '10.7' and r[6] == 'NC' for r in lignes))
        j = telecharger(page, 'export-jira')
        rows = list(csv.reader(io.StringIO(j.read_text(encoding='utf-8-sig'))))
        menu = [r for r in rows[1:] if 'focus invisible' in r[0]]
        verifier('Jira : un seul ticket pour l\'observation valable sur 3 écrans', len(menu) == 1, [r[0] for r in rows[1:]])
        verifier('Jira : référence RAWeb', 'RAWeb 1.1, Criterion 10.7' in menu[0][1] if menu else False)
        verifier('Jira : observation retestée sur un seul écran toujours exporté (encore NC ailleurs)', len(menu) == 1)
        s = telecharger(page, 'export-json')
        sauvegarde = json.loads(s.read_text(encoding='utf-8'))
        verifier('Sauvegarde : version 2', sauvegarde.get('v') == 2)

        # Import : sauvegarde v2 relue, v1 refusée avec un message.
        page.click('#btn-menu')
        page.click('[data-action="reinit"]')
        page.click('[data-action="reinit"]')
        verifier('Nouvel audit : vide', page.evaluate('RUXAWEB.state().onglets.length') == 0)
        page.set_input_files('#fichier-import', str(s))
        page.wait_for_timeout(300)
        verifier('Import v2 : onglet restauré', page.evaluate('RUXAWEB.state().onglets.length') == 1)
        v1 = pathlib.Path(tempfile.mkdtemp()) / 'v1.json'
        v1.write_text(json.dumps({'onglets': [], 'auditNom': 'ancien'}), encoding='utf-8')
        page.set_input_files('#fichier-import', str(v1))
        page.wait_for_timeout(300)
        verifier('Import v1 : refusé avec renvoi vers 1.7', '1.7' in page.locator('#statut').inner_text())

        # Scénario : transition entre deux écrans.
        page.click('[data-action="accueil"]')
        page.click('[data-action="nouvel-onglet"]')
        page.check('input[name="onglet-type"][value="scenario"]')
        page.click('[data-action="vue"][data-vue="audit"] >> nth=-1')
        verifier('Scénario : champ Transition vers', page.locator('#obs-vers').count() == 1)
        page.fill('#obs-texte', 'Le focus reste sur le bouton Continuer après le passage à l\'étape suivante.')
        page.select_option('#obs-point', 'ordre-focus')
        page.select_option('#obs-vers', index=1)
        page.click('[data-action="tag"][data-tag="nc"]')
        page.click('#form-obs button[type=submit]')
        verifier('Scénario : transition affichée Écran -> Écran', ' -> ' in page.locator('#t-obs').locator('xpath=..').inner_text())
        page.click('[data-action="vue"][data-vue="ux"]')
        verifier('Scénario : onglet POUR/UX avec 8 questions', page.locator('[data-action="heur"][data-v="succes"]').count() == 8)
        for i in range(6):
            page.locator('[data-action="heur"][data-v="succes"]').nth(i).click()
        page.locator('[data-action="heur"][data-v="echec"]').nth(6).click()
        page.locator('[data-action="heur"][data-v="na"]').nth(7).click()
        page.fill('textarea[data-heur-com="h7"]', 'Six étapes pour une simple commande.')
        page.locator('textarea[data-heur-com="h7"]').blur()
        oid2 = page.evaluate('RUXAWEB.state().ongletId')
        verifier('POUR/UX : score 86 % (6 succès sur 7), hors taux', '86 %' in page.locator('main').inner_text())
        verifier('POUR/UX : commentaire gardé', page.evaluate('RUXAWEB.state().onglets.filter(function(o){return o.id==="%s"})[0].heuristiques.h7.commentaire' % oid2) == 'Six étapes pour une simple commande.')
        page.click('[data-action="vue"][data-vue="resume"]')
        verifier('Résumé du scénario : score POUR/UX à part', 'Score POUR/UX' in page.locator('main').inner_text())
        x2 = telecharger(page, 'export-xlsx')
        verifier('Excel : lignes POUR/UX du scénario', sum(1 for r in openpyxl.load_workbook(x2).active.iter_rows(values_only=True) if str(r[1] or '').startswith('POUR/UX')) == 8)
        jira = telecharger(page, 'export-jira')
        verifier('Scénario : ticket avec la transition', any(' -> ' in r[0] for r in csv.reader(io.StringIO(jira.read_text(encoding='utf-8-sig')))))
        page.click('[data-action="accueil"]')
        page.locator('[data-action="ouvrir"]').first.click()

        # Import d'un rapport axe-core puis Lighthouse : des résultats à reprendre, jamais un statut.
        page.click('[data-action="vue"][data-vue="audit"]')
        ecran(page, 1)
        axe = {'testEngine': {'name': 'axe-core', 'version': '4.13.0'}, 'url': 'https://exemple.lu/',
               'violations': [{'id': 'image-alt', 'help': 'Images must have alternative text', 'tags': ['wcag2a', 'wcag111'],
                               'nodes': [{'target': ['#logo']}, {'target': ['.promo img']}]},
                              {'id': 'regle-inconnue', 'help': 'Contraste insuffisant (test)', 'tags': ['wcag2aa', 'wcag143'],
                               'nodes': [{'target': ['p.note']}]}],
               'incomplete': [{'id': 'color-contrast', 'help': 'Contrast to review', 'tags': ['wcag143'], 'nodes': [{'target': ['h2']}]}],
               'passes': [{'id': 'document-title', 'tags': ['wcag242'], 'nodes': [{'target': ['html']}]}]}
        f = pathlib.Path(tempfile.mkdtemp()) / 'axe.json'
        f.write_text(json.dumps(axe), encoding='utf-8')
        avant = statut(page, 1, 'raweb-1-1')
        page.click('[data-action="vue"][data-vue="audit"]')
        page.set_input_files('#fichier-auto', str(f))
        page.wait_for_timeout(300)
        verifier('axe : 3 résultats à vérifier (2 échecs, 1 incomplet)', page.locator('[data-action="reprendre"]').count() == 3)
        verifier('axe : règle réussie ignorée', 'document-title' not in page.locator('main').inner_text())
        verifier('axe : rien de compté avant reprise', statut(page, 1, 'raweb-1-1') == avant)
        cands = page.evaluate('RUXAWEB.state().onglets[0].candidats')
        par = {c['regle']: c['criteres'] for c in cands}
        verifier('axe : image-alt vers 1.1 (table RUXAL)', par.get('image-alt') == ['raweb-1-1'], par)
        verifier('axe : règle hors table rapprochée par WCAG 1.4.3 (3.2)', par.get('regle-inconnue') == ['raweb-3-2'], par)
        page.click('[data-action="vue"][data-vue="audit"]')
        page.locator('[data-action="reprendre"]').first.click()
        verifier('Reprendre : observation préremplie', 'Images must have alternative text' in page.input_value('#obs-texte'))
        page.click('#form-obs button[type=submit]')
        derniere = page.evaluate('RUXAWEB.state().onglets[0].obs.slice(-1)[0]')
        verifier('Reprendre : observation #nc avec sa source', derniere['tag'] == 'nc' and derniere.get('source', {}).get('regle') == 'image-alt', derniere)
        verifier('Reprendre : 1.1 NC', statut(page, 1, 'raweb-1-1') == 'NC')
        page.click('[data-action="vue"][data-vue="audit"]')
        verifier('Reprendre : résultat retiré de la liste', page.locator('[data-action="reprendre"]').count() == 2)
        page.locator('[data-action="ecarter"]').first.click()
        verifier('Écarter : résultat retiré', page.locator('[data-action="reprendre"]').count() == 1)
        lh = {'lighthouseVersion': '13.5.0', 'categories': {'accessibility': {'auditRefs': [{'id': 'button-name'}, {'id': 'html-has-lang'}, {'id': 'focus-traps'}]}},
              'audits': {'button-name': {'id': 'button-name', 'title': 'Buttons do not have an accessible name', 'score': 0, 'scoreDisplayMode': 'binary',
                                         'details': {'items': [{'node': {'selector': 'button.close'}}]}},
                         'html-has-lang': {'id': 'html-has-lang', 'title': 'html has lang', 'score': 1, 'scoreDisplayMode': 'binary'},
                         'focus-traps': {'id': 'focus-traps', 'title': 'Manual', 'score': None, 'scoreDisplayMode': 'manual'}}}
        f2 = f.parent / 'lh.json'
        f2.write_text(json.dumps(lh), encoding='utf-8')
        page.set_input_files('#fichier-auto', str(f2))
        page.wait_for_timeout(300)
        verifier('Lighthouse : un seul échec, réussites et contrôles manuels ignorés', page.locator('[data-action="reprendre"]').count() == 1 and 'Buttons do not have' in page.locator('main').inner_text())
        cands = [c for c in page.evaluate('RUXAWEB.state().onglets[0].candidats') if c['etat'] == 'ouvert']
        verifier('Lighthouse : button-name vers 7.1 et 11.9', cands and cands[0]['criteres'] == ['raweb-7-1', 'raweb-11-9'], cands)
        f3 = f.parent / 'autre.json'
        f3.write_text(json.dumps({'foo': 1}), encoding='utf-8')
        page.set_input_files('#fichier-auto', str(f3))
        page.wait_for_timeout(300)
        verifier('Import d\'un autre fichier : refusé avec un message', 'ni un rapport axe' in page.locator('#statut').inner_text())

        page.emulate_media(color_scheme='dark')
        page.click('[data-action="accueil"]')
        verifier('Mode sombre : pas d\'erreur', not erreurs, erreurs)
        verifier('Aucune erreur JavaScript', not erreurs, erreurs)
        nav.close()
    echecs = resultats.count(False)
    print('\n%d contrôles, %d échec(s)' % (len(resultats), echecs))
    return 1 if echecs else 0


if __name__ == '__main__':
    sys.exit(main())
